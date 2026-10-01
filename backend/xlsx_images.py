"""
xlsx_images.py — Excel çıktılarına ÜRÜN GÖRSELİ gömme (tek ortak altyapı).

Kural (kullanıcı isteği): görseller TAM görünecek — üstten/alttan kırpma YOK, hepsi aynı
ölçüde ve hücreye tam oturacak; rapor patrona sunulacak kalitede görünmeli.

Yöntem:
  * Görsel oranı KORUNARAK (contain) kutuya sığdırılır → kırpma yok.
  * Sonuç, sabit ölçülü BEYAZ tuvale ORTALANIR → her hücrede birebir aynı boyut,
    sütun hizası bozulmaz (moda fotoğrafları dikey, kimi ürünler kare/yatay olabilir).
  * Satır yüksekliği ve sütun genişliği tuvale göre ayarlanır → taşma/kırpılma olmaz.
"""
from __future__ import annotations

import asyncio
import logging
from io import BytesIO
from typing import Dict, Iterable, Optional

logger = logging.getLogger(__name__)

# (genişlik, yükseklik) piksel — moda fotoğrafı dikey olduğundan 3:4 tuval
SIZES = {
    "small":  (96, 128),
    "medium": (132, 176),
    "large":  (180, 240),
}
_PAD_PX = 4          # tuval içi boşluk (görsel hücre kenarına yapışmasın)
_MAX_IMAGES = 1500   # tek dosyada indirilecek en fazla BENZERSİZ görsel
_TIME_BUDGET = 150   # saniye — indirme toplam süresi (istek zaman aşımına düşmesin)


def box_for(size: str) -> tuple:
    return SIZES.get(str(size or "medium").lower(), SIZES["medium"])


def row_height_pt(size: str) -> float:
    """Tuvalin tam sığacağı satır yüksekliği (punto). 1 px = 0.75 pt."""
    return round(box_for(size)[1] * 0.75 + 3, 1)


def col_width_chars(size: str) -> float:
    """Tuvalin tam sığacağı sütun genişliği (Excel karakter birimi ≈ px/7)."""
    return round(box_for(size)[0] / 7.0 + 1.2, 1)


def fit_canvas(raw: bytes, size: str = "medium") -> Optional[bytes]:
    """Ham görseli, oranını KORUYARAK sabit tuvale ortalar. Kırpma yapmaz."""
    try:
        from PIL import Image, ImageOps
    except Exception:
        return None
    try:
        bw, bh = box_for(size)
        im = Image.open(BytesIO(raw))
        # Şeffaf PNG/WebP → beyaz zemin (Excel'de siyah blok görünmesin)
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA")
            bg = Image.new("RGB", im.size, (255, 255, 255))
            bg.paste(im, mask=im.split()[-1])
            im = bg
        else:
            im = im.convert("RGB")
        inner = (max(1, bw - 2 * _PAD_PX), max(1, bh - 2 * _PAD_PX))
        im = ImageOps.contain(im, inner, Image.LANCZOS)   # KIRPMA YOK
        canvas = Image.new("RGB", (bw, bh), (255, 255, 255))
        canvas.paste(im, ((bw - im.width) // 2, (bh - im.height) // 2))
        out = BytesIO()
        # JPEG q92: gerçek ürün fotoğrafında PNG'nin ~1/4'ü boyut (17 KB → 4.8 KB) ve
        # dolgu beyazı bozulmaz (kenar parlaklığı 254.6/255). 1000 ürünlü raporda dosya
        # 17 MB yerine ~5 MB kalır — e-postayla gönderilebilir, Excel hızlı açar.
        canvas.save(out, format="JPEG", quality=92, optimize=True)
        return out.getvalue()
    except Exception as e:
        logger.debug(f"[xlsx-img] görsel işlenemedi: {type(e).__name__}: {e}")
        return None


async def fetch_images(urls: Iterable[str], size: str = "medium",
                       concurrency: int = 8) -> Dict[str, bytes]:
    """Benzersiz URL'leri eşzamanlı indirip tuvale oturtur. Hata → o görsel atlanır."""
    uniq, seen = [], set()
    for u in urls:
        u = str(u or "").strip()
        if u and u not in seen and u.startswith(("http://", "https://")):
            seen.add(u)
            uniq.append(u)
        if len(uniq) >= _MAX_IMAGES:
            break
    if not uniq:
        return {}
    try:
        import httpx
    except Exception:
        return {}
    out: Dict[str, bytes] = {}
    sem = asyncio.Semaphore(max(1, concurrency))
    headers = {"User-Agent": "Mozilla/5.0 (StoreExcel/1.0)"}

    async def grab(cli, u):
        async with sem:
            try:
                r = await cli.get(u)
                if r.status_code == 200 and r.content:
                    png = await asyncio.to_thread(fit_canvas, r.content, size)  # PIL → olay döngüsünü kilitlemesin
                    if png:
                        out[u] = png
            except Exception:
                pass

    try:
        async with httpx.AsyncClient(timeout=12.0, follow_redirects=True, headers=headers) as cli:
            await asyncio.wait_for(
                asyncio.gather(*[grab(cli, u) for u in uniq]), timeout=_TIME_BUDGET)
    except asyncio.TimeoutError:
        logger.warning(f"[xlsx-img] süre sınırı: {len(out)}/{len(uniq)} görsel indirildi")
    except Exception as e:
        logger.warning(f"[xlsx-img] indirme hatası: {e}")
    logger.info(f"[xlsx-img] {len(out)}/{len(uniq)} görsel hazır (boyut={size})")
    return out


def first_image_url(p: dict) -> str:
    """Üründen İLK GERÇEK görsel (beden tablosu vb. atlanır)."""
    for im in (p.get("images") or []):
        if isinstance(im, str) and im.strip():
            return im.strip()
        if isinstance(im, dict) and not im.get("is_size_table"):
            u = im.get("url") or im.get("src") or im.get("image")
            if u:
                return str(u).strip()
    for k in ("image", "thumbnail", "cover_image", "main_image"):
        if p.get(k):
            return str(p[k]).strip()
    return ""


def prepare_sheet(ws, col_idx: int, size: str = "medium", header_row: int = 1) -> None:
    """Görsel sütununun genişliğini ayarlar (satır yüksekliği satır bazında verilir)."""
    from openpyxl.utils import get_column_letter
    ws.column_dimensions[get_column_letter(col_idx)].width = col_width_chars(size)


def put_image(ws, row_idx: int, col_idx: int, png: bytes, size: str = "medium") -> bool:
    """Hücreye görseli yerleştirir + satır yüksekliğini TAM sığacak şekilde ayarlar."""
    if not png:
        return False
    try:
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.utils import get_column_letter
        img = XLImage(BytesIO(png))
        bw, bh = box_for(size)
        img.width, img.height = bw, bh          # birebir tuval ölçüsü (esneme yok)
        ws.add_image(img, f"{get_column_letter(col_idx)}{row_idx}")
        _need = row_height_pt(size)
        _cur = ws.row_dimensions[row_idx].height or 0
        if _cur < _need:
            ws.row_dimensions[row_idx].height = _need
        return True
    except Exception as e:
        logger.debug(f"[xlsx-img] hücreye eklenemedi: {e}")
        return False
