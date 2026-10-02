"""routes/report_assistant.py — Raporlar sayfasının "Bana Sor" asistanı.

Yönetici doğal dille (yazarak ya da sesli) rapor sorusu sorar; dil modeli sistemin
KENDİ rapor fonksiyonlarını araç olarak çağırarak cevap verir. Model rakam UYDURMAZ:
her sayı, raporlar ekranının kullandığı aynı hesaplardan (dönem muhasebesi, kohort
net, gider pusulası = iade, TR yerel gün sınırları, mükerrer/ticimax elemeleri) gelir.

Sağlayıcı: önce Gemini, olmazsa (anahtar yok / kota / hata) OpenAI.

⛔ Gizlilik: araçlar dil modeline MÜŞTERİ KİŞİSEL VERİSİ göndermez — ad, telefon,
e-posta, adres, TC, IBAN hiçbir araç çıktısında yer almaz. Yalnız toplamlar, ürün,
sipariş numarası, statü, tutar, tarih, il/ilçe (toplu) gider.
"""

import contextvars
import inspect
import io
import json
import uuid
import os
import re
import time
from datetime import datetime, timedelta, timezone
from typing import Any, Optional

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import Response

from .deps import db, limiter, logger, require_admin

router = APIRouter(prefix="/admin/reports/assistant", tags=["report-assistant"])

_TR = timedelta(hours=3)
_KANALLAR = ["all", "site", "trendyol", "hepsiburada", "amazon", "temu"]
_MAX_TUR = 6            # model ↔ araç gidiş-dönüş sınırı (sonsuz döngü olmasın)
_MAX_SONUC = 14000      # tek araç çıktısının modele giden en fazla karakteri
_MAX_SES = 10 * 1024 * 1024


# ── Anahtar + model çözümü ────────────────────────────────────────────────────
async def _anahtar(ad: str, saglayici: str) -> str:
    """Önce ortam değişkeni, sonra şifreli kasa."""
    v = (os.environ.get(ad) or "").strip()
    if v:
        return v
    try:
        from .secrets_vault import get_secret
        v = (await get_secret(ad) or "").strip()
        if v:
            return v
    except Exception:
        pass
    return ""


async def _ayarlar() -> dict:
    s = await db.settings.find_one({"id": "report_assistant"}, {"_id": 0}) or {}
    return {
        "enabled": s.get("enabled", True),
        "gemini_model": s.get("gemini_model") or "gemini-3.1-flash",
        "openai_model": s.get("openai_model") or "gpt-5.4-mini",
    }


# ── Tarih yardımcıları ────────────────────────────────────────────────────────
def _tr_simdi() -> datetime:
    return datetime.now(timezone.utc) + _TR


def _gun(v: Any, varsayilan: str) -> str:
    s = str(v or "").strip()[:10]
    try:
        datetime.strptime(s, "%Y-%m-%d")
        return s
    except Exception:
        return varsayilan


def _aralik(a: dict) -> tuple[str, str]:
    now = _tr_simdi().date()
    bas = _gun(a.get("start_date"), now.replace(day=1).isoformat())
    bit = _gun(a.get("end_date"), now.isoformat())
    if bas > bit:
        bas, bit = bit, bas
    return bas, bit


def _kanal(v: Any) -> Optional[str]:
    k = str(v or "all").strip().lower()
    return None if k in ("", "all", "tum", "tümü", "hepsi") else (k if k in _KANALLAR else None)


# ── Rapor fonksiyonunu doğrudan çağır (FastAPI Query varsayılanlarını çöz) ────
async def _cagir(fn, user: dict, **kw):
    """Uç fonksiyonları doğrudan çağrılınca `Query(None)` varsayılanı None DEĞİL,
    bir FieldInfo nesnesi olarak gelir ve süzgeçleri bozar. Verilmeyen her parametre
    için Query'nin gerçek varsayılanı kullanılır."""
    args = {}
    for ad, p in inspect.signature(fn).parameters.items():
        if ad == "current_user":
            args[ad] = user
        elif ad in kw:
            args[ad] = kw[ad]
        elif p.default is not inspect.Parameter.empty:
            d = p.default
            args[ad] = getattr(d, "default", d) if type(d).__module__.startswith("fastapi") else d
    return await fn(**args)


def _yuvarla(v):
    try:
        return round(float(v or 0), 2)
    except Exception:
        return 0.0


def _kova(d: dict) -> dict:
    d = d or {}
    return {"tutar": _yuvarla(d.get("revenue")), "siparis": int(d.get("orders") or 0),
            "adet": int(d.get("units") or 0)}


# ── ARAÇLAR ───────────────────────────────────────────────────────────────────
async def _t_donem_ozeti(a, u):
    from . import reports as R
    s, e = _aralik(a)
    b = await _cagir(R.sales_breakdown, u, start_date=s, end_date=e, source=_kanal(a.get("kanal")))
    od = b.get("onceki_donem") or {}
    return {
        "aralik": [s, e], "kanal": _kanal(a.get("kanal")) or "tüm kanallar",
        "brut_ciro_iptal_iade_dahil": _kova(b.get("included")),
        "iptal_bu_donemde_kesinlesen": _kova(b.get("cancels")),
        "iade_bu_donemde_pusulasi_kesilen": _kova(b.get("returns")),
        "net_DONEM_muhasebe": _kova(b.get("net")),
        "net_BU_DONEMIN_SATISININ_kohort": _kova(b.get("net_kohort")),
        "onceki_aylarin_siparislerine_ait_iade": _kova((od.get("iade") or {})),
        "onceki_aylarin_siparislerine_ait_iptal": _kova((od.get("iptal") or {})),
        "acik_iade_talepleri": _kova(b.get("pending_returns")),
        "mukerrer_elenen_pusula": b.get("mukerrer_elenen"),
    }


async def _t_kanal_tablosu(a, u):
    from . import reports as R
    s, e = _aralik(a)
    r = await _cagir(R.cancel_return_by_source, u, start_date=s, end_date=e)
    out = []
    for it in (r.get("items") or []):
        out.append({
            "kanal": it.get("source"),
            "brut": {"tutar": _yuvarla(it.get("total_revenue")), "siparis": it.get("total_orders"),
                     "adet": it.get("total_units")},
            "iptal_bu_donemin_satisindan": {"tutar": _yuvarla(it.get("cancel_total")),
                                            "siparis": it.get("cancel_orders")},
            "iade_bu_donemin_satisindan": {"tutar": _yuvarla(it.get("return_total")),
                                           "siparis": it.get("return_orders"),
                                           "adet": it.get("return_units")},
            "net_bu_donemin_satisi": {"tutar": _yuvarla(it.get("revenue")),
                                      "siparis": it.get("orders"), "adet": it.get("units")},
            "net_donem_muhasebe": _yuvarla(it.get("donem_revenue")),
            "acik_iade": {"tutar": _yuvarla(it.get("pending_total")), "adet": it.get("pending_units")},
        })
    return {"aralik": [s, e], "kanallar": out}


async def _t_zaman_serisi(a, u):
    from . import reports as R
    s, e = _aralik(a)
    g = str(a.get("gruplama") or "day").lower()
    g = {"gun": "day", "gün": "day", "hafta": "week", "ay": "month"}.get(g, g)
    g = g if g in ("day", "week", "month") else "day"
    r = await _cagir(R.sales, u, start_date=s, end_date=e, group_by=g, source=_kanal(a.get("kanal")))
    rows = r.get("rows") or []
    return {"aralik": [s, e], "gruplama": g, "satirlar": rows[:120], "toplam": r.get("totals")}


_URUN_ALANLARI = ("name", "stock_code", "category", "category_name", "qty", "revenue",
                  "gross_revenue", "discount_amount", "cancel_qty", "cancel_total", "return_qty",
                  "return_total", "current_stock", "best_size", "top_platform",
                  "return_rate_excluding_cancels_pct", "trendyol_return_rate_pct")


async def _t_urun_raporu(a, u):
    from . import reports as R
    s, e = _aralik(a)
    r = await _cagir(R.top_products, u, limit=5000, start_date=s, end_date=e,
                     source=_kanal(a.get("kanal")))
    items = r.get("items") or []
    q = str(a.get("arama") or "").strip().lower()
    if q:
        items = [i for i in items if q in str(i.get("name") or "").lower()
                 or q in str(i.get("stock_code") or "").lower()
                 or q in str(i.get("barcode") or "").lower()]
    sirala = str(a.get("siralama") or "ciro").lower()
    key = {"ciro": "revenue", "adet": "qty", "iade": "return_qty", "iptal": "cancel_qty",
           "stok": "current_stock"}.get(sirala, "revenue")
    items = sorted(items, key=lambda i: -float(i.get(key) or 0))
    lim = max(1, min(int(a.get("limit") or (5000 if a.get("_excel") else 15)), 5000 if a.get("_excel") else 50))
    out = []
    for i in items[:lim]:
        row = {k: i.get(k) for k in _URUN_ALANLARI if i.get(k) not in (None, "")}
        v = i.get("velocity")
        if isinstance(v, dict):
            row["satis_hizi"] = v.get("label")
        row["beden_kirilimi"] = (i.get("size_breakdown") or [])[:8]
        row["kanal_kirilimi"] = (i.get("platform_breakdown") or [])[:6]
        out.append(row)
    return {"aralik": [s, e], "eslesen_urun": len(items), "siralama": sirala, "urunler": out}


async def _t_kategori(a, u):
    from . import reports as R
    s, e = _aralik(a)
    r = await _cagir(R.category_report, u, start_date=s, end_date=e, source=_kanal(a.get("kanal")))
    return {"aralik": [s, e], "kategoriler": (r.get("items") or [])[:40]}


async def _t_odeme(a, u):
    from . import reports as R
    s, e = _aralik(a)
    r = await _cagir(R.payment_report, u, start_date=s, end_date=e, source=_kanal(a.get("kanal")))
    return {"aralik": [s, e], "odeme_yontemleri": r.get("items") or []}


async def _t_saat_gun(a, u):
    from . import reports as R
    s, e = _aralik(a)
    tur = str(a.get("tur") or "saat").lower()
    fn = R.sales_by_weekday if tur.startswith("g") else R.sales_by_hour
    r = await _cagir(fn, u, start_date=s, end_date=e, source=_kanal(a.get("kanal")))
    return {"aralik": [s, e], "tur": "gün" if tur.startswith("g") else "saat (TR)",
            "satirlar": r.get("rows"), "zirve": r.get("peak")}


async def _t_konum(a, u):
    from . import reports as R
    s, e = _aralik(a)
    grp = "district" if str(a.get("seviye") or "").lower() in ("ilce", "ilçe", "district") else "city"
    lim = max(1, min(int(a.get("limit") or (1000 if a.get("_excel") else 15)), 1000 if a.get("_excel") else 60))
    r = await _cagir(R.sales_by_location, u, start_date=s, end_date=e, group=grp,
                     source=_kanal(a.get("kanal")), limit=lim)
    return {"aralik": [s, e], "seviye": grp, "satirlar": r.get("rows"), "toplam": r.get("totals")}


async def _t_iade_nedenleri(a, u):
    from . import reports as R
    s, e = _aralik(a)
    r = await _cagir(R.returns_by_reason, u, start_date=s, end_date=e, source=_kanal(a.get("kanal")))
    return {"aralik": [s, e], "nedenler": (r.get("reasons") or [])[:(1000 if a.get("_excel") else 30)]}


async def _t_iptal_iade_urunleri(a, u):
    from . import reports as R
    s, e = _aralik(a)
    r = await _cagir(R.cancel_return_products, u, start_date=s, end_date=e,
                     source=_kanal(a.get("kanal")))
    rows = r.get("items") or r.get("rows") or []
    key = "cancel_qty" if str(a.get("tur") or "").lower().startswith("ipt") else "return_qty"
    rows = sorted(rows, key=lambda x: -float(x.get(key) or 0))
    lim = max(1, min(int(a.get("limit") or (5000 if a.get("_excel") else 15)), 5000 if a.get("_excel") else 50))
    return {"aralik": [s, e], "siralama": key, "urunler": rows[:lim]}


async def _t_stok(a, u):
    from . import reports as R
    r = await _cagir(R.stock_report, u)
    kisa = {k: v for k, v in r.items() if not isinstance(v, list)}
    for k, v in r.items():
        if isinstance(v, list):
            kisa[k + "_adet"] = len(v)
            kisa[k + ("_liste" if a.get("_excel") else "_ilk30")] = v if a.get("_excel") else v[:30]
    return kisa


async def _t_hizli_satan(a, u):
    from . import reports as R
    gun = max(1, min(int(a.get("gun") or 14), 90))
    lim = max(1, min(int(a.get("limit") or (500 if a.get("_excel") else 20)), 500 if a.get("_excel") else 60))
    r = await _cagir(R.fast_selling_products, u, window_days=gun, min_sold=1, limit=lim)
    return {"son_gun": gun, "urunler": r.get("items")}


async def _t_kupon(a, u):
    from . import reports as R
    s, e = _aralik(a)
    r = await _cagir(R.coupon_performance, u, start_date=s, end_date=e)
    return {"aralik": [s, e], **r}


async def _t_gunun_urunleri(a, u):
    from . import reports as R
    s, e = _aralik(a)
    r = await _cagir(R.day_orders, u, start_date=s, end_date=e, source=_kanal(a.get("kanal")))
    return {"aralik": [s, e], "siparis_sayisi": r.get("order_count"),
            "toplam_adet": r.get("total_qty"),
            "urun_beden": (r.get("rows") or [])[:(100000 if a.get("_excel") else 60)]}


_STATU_TR = {
    "confirmed": "Onaylandı", "preparing": "Hazırlanıyor", "shipped": "Kargoda",
    "delivered": "Teslim edildi", "cancelled": "İptal", "cancel_refunded": "İptal (iade edildi)",
    "return_requested": "İade talebi", "return_in_transit": "İade kargoda",
    "return_approved": "İade onaylandı", "returned": "İade edildi",
    "refunded": "İade bedeli ödendi", "partial_refunded": "Kısmi iade",
    "awaiting_payment": "Ödeme bekleniyor", "payment_failed": "Ödeme başarısız",
    "pending": "Beklemede",
}


def _tr_saat(iso: Any) -> str:
    try:
        d = datetime.fromisoformat(str(iso).replace("Z", "+00:00"))
        if d.tzinfo is None:
            d = d.replace(tzinfo=timezone.utc)
        return (d.astimezone(timezone.utc) + _TR).strftime("%d.%m.%Y %H:%M")
    except Exception:
        return ""


async def _t_siparis(a, u):
    no = str(a.get("siparis_no") or "").strip()
    if not no:
        return {"hata": "siparis_no gerekli"}
    o = await db.orders.find_one(
        {"order_number": no},
        {"_id": 0, "order_number": 1, "platform": 1, "status": 1, "total": 1, "subtotal": 1,
         "discount_amount": 1, "shipping_cost": 1, "payment_method": 1, "payment_status": 1,
         "marketplace_order_date": 1, "created_at": 1, "cancelled_at": 1,
         "partial_cancel_amount": 1, "cargo_provider_name": 1, "cargo_tracking_number": 1,
         "items.name": 1, "items.product_name": 1, "items.size": 1, "items.color": 1,
         "items.quantity": 1, "items.price": 1, "items.barcode": 1})
    if not o:
        return {"bulunamadi": no}
    kalemler = [{"urun": i.get("name") or i.get("product_name"), "beden": i.get("size"),
                 "renk": i.get("color"), "adet": i.get("quantity"), "fiyat": i.get("price"),
                 "barkod": i.get("barcode")} for i in (o.get("items") or [])]
    pus = [{"kocan_no": g.get("number") or g.get("display_number"),
            "tarih": _tr_saat(g.get("created_at")), "tutar": (g.get("totals") or {}).get("net")}
           async for g in db.gider_pusulasi.find(
               {"order_number": no},
               {"_id": 0, "number": 1, "display_number": 1, "created_at": 1, "totals.net": 1})]
    talep = [{"statu": c.get("claim_status"), "acilis": _tr_saat(c.get("created_date")),
              "onay": _tr_saat(c.get("return_approved_at"))}
             async for c in db.trendyol_claims.find(
                 {"order_number": no},
                 {"_id": 0, "claim_status": 1, "created_date": 1, "return_approved_at": 1})]
    return {
        "siparis_no": no, "kanal": o.get("platform") or "site",
        "statu": _STATU_TR.get(str(o.get("status")), o.get("status")),
        "siparis_tarihi_TR": _tr_saat(o.get("marketplace_order_date") or o.get("created_at")),
        "iptal_tarihi_TR": _tr_saat(o.get("cancelled_at")) if o.get("cancelled_at") else None,
        "tutar": o.get("total"), "ara_toplam": o.get("subtotal"),
        "indirim": o.get("discount_amount"), "kargo_ucreti": o.get("shipping_cost"),
        "kismi_iptal_tutari": o.get("partial_cancel_amount"),
        "odeme": o.get("payment_method"), "odeme_durumu": o.get("payment_status"),
        "kargo": {"firma": o.get("cargo_provider_name"), "takip": o.get("cargo_tracking_number")},
        "kalemler": kalemler, "gider_pusulalari": pus, "iade_talepleri": talep,
    }


async def _t_siparis_listesi(a, u):
    """Dönemin sipariş listesi — raporlar ekranıyla AYNI küme (_sales_stages +
    dönem hariç tutması: ödenmemişler yok). Müşteri kişisel verisi YOK."""
    from . import reports as R
    s, e = _aralik(a)
    ss, ee = R._iso_range(s, e)
    lim = 20000 if a.get("_excel") else max(1, min(int(a.get("limit") or 30), 50))
    statu = str(a.get("statu") or "").strip().lower()
    pipe = [*R._sales_stages(ss, ee, _kanal(a.get("kanal")), exclude=R._DONEM_EXCLUDED)]
    if statu == "iptal":
        pipe.append({"$match": {"status": {"$in": ["cancelled", "cancel_refunded"]}}})
    elif statu == "aktif":
        pipe.append({"$match": {"status": {"$nin": ["cancelled", "cancel_refunded"]}}})
    pipe += [{"$sort": {"marketplace_order_date": 1, "created_at": 1}}, {"$limit": lim},
             {"$project": {"_id": 0, "order_number": 1, "platform": 1, "status": 1, "total": 1,
                           "marketplace_order_date": 1, "created_at": 1, "cancelled_at": 1,
                           "payment_method": 1, "items.name": 1, "items.product_name": 1,
                           "items.size": 1, "items.quantity": 1, "items.barcode": 1}}]
    rows, toplam, n = [], 0.0, 0
    async for o in db.orders.aggregate(pipe):
        its = o.get("items") or []
        adet = sum(int(i.get("quantity") or 1) for i in its)
        t = R._fnum(o.get("total"))
        toplam += t
        n += 1
        rows.append({
            "siparis_no": o.get("order_number"), "kanal": o.get("platform") or "site",
            "tarih_TR": _tr_saat(o.get("marketplace_order_date") or o.get("created_at")),
            "statu": _STATU_TR.get(str(o.get("status")), o.get("status")),
            "tutar": round(t, 2), "adet": adet, "odeme": o.get("payment_method"),
            "iptal_tarihi_TR": _tr_saat(o.get("cancelled_at")) if o.get("cancelled_at") else "",
            "urunler": "; ".join(f"{i.get('product_name') or i.get('name') or 'Ürün'}"
                                 f"{(' ' + str(i.get('size'))) if i.get('size') else ''}"
                                 f" x{int(i.get('quantity') or 1)}" for i in its),
        })
    return {"aralik": [s, e], "listelenen": n, "listelenen_tutar": round(toplam, 2),
            "not": "Ödenmemiş (ödeme bekleyen/başarısız) siparişler listede yoktur.",
            "siparisler": rows}


# ── Excel ────────────────────────────────────────────────────────────────────
# Sohbet cevabıyla birlikte dönecek dosyalar (istek başına; araç katmanından yazılır).
_DOSYALAR: contextvars.ContextVar = contextvars.ContextVar("rapor_asistan_dosyalar", default=None)

_BASLIK_TR = {
    "name": "Ürün", "product_name": "Ürün", "size": "Beden", "color": "Renk", "barcode": "Barkod",
    "stock_code": "Stok kodu", "qty": "Adet", "quantity": "Adet", "revenue": "Ciro (TL)",
    "orders": "Sipariş", "units": "Ürün adedi", "return_qty": "İade adedi",
    "cancel_qty": "İptal adedi", "current_stock": "Stok", "partial_orders": "Kısmi sipariş",
    "city": "İl", "district": "İlçe", "count": "Adet", "total": "Tutar (TL)",
    "reason": "Neden", "label": "Etiket", "date": "Tarih", "channel": "Kanal",
    "platform": "Kanal", "category": "Kategori", "avg": "Ortalama",
}


def _hucre(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Evet" if v else "Hayır"
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, dict):
        return ", ".join(f"{k}: {_hucre(x)}" for k, x in v.items())
    if isinstance(v, list):
        return "; ".join(str(_hucre(x)) for x in v)
    return str(v)


def _duz(d: dict, on: str = "") -> dict:
    out = {}
    for k, v in (d or {}).items():
        key = f"{on}{k}"
        if isinstance(v, dict) and v and not any(isinstance(x, (dict, list)) for x in v.values()) and len(v) <= 8:
            out.update(_duz(v, key + " · "))
        else:
            out[key] = v
    return out


def _sayfalara_bol(res: dict) -> list:
    """Araç çıktısını sayfalara çevirir: satır listeleri → ayrı sayfa, tekil değerler → Özet."""
    sayfalar, ozet = [], []
    for k, v in (res or {}).items():
        if isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
            sayfalar.append((k, [_duz(x) for x in v]))
        elif isinstance(v, dict) and v and all(isinstance(x, dict) for x in v.values()):
            sayfalar.append((k, [{"_": kk, **_duz(vv)} for kk, vv in v.items()]))
        elif isinstance(v, dict):
            ozet += [(f"{k} · {kk}", vv) for kk, vv in _duz(v).items()]
        else:
            ozet.append((k, v))
    if ozet:
        sayfalar.insert(0, ("Özet", [{"Alan": a, "Değer": b} for a, b in ozet]))
    return sayfalar


def _xlsx(baslik: str, sayfalar: list) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    kullanilan = set()
    for ad, rows in sayfalar:
        sad = re.sub(r"[\[\]\:\*\?\/\\]", " ", str(ad).replace("_", " ")).strip()[:31] or "Sayfa"
        while sad in kullanilan:
            sad = (sad[:28] + f" {len(kullanilan)}")[:31]
        kullanilan.add(sad)
        ws = wb.create_sheet(sad)
        cols: list = []
        for r in rows:
            for k in r:
                if k not in cols:
                    cols.append(k)
        ws.append([_BASLIK_TR.get(c, str(c).replace("_", " ")) if c != "_" else "" for c in cols])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="111827")
            c.alignment = Alignment(vertical="center")
        for r in rows[:50000]:
            ws.append([_hucre(r.get(c)) for c in cols])
        for i, c in enumerate(cols, start=1):
            harf = get_column_letter(i)
            uz = max([len(str(ws.cell(row=1, column=i).value or ""))] +
                     [len(str(x.value)) for x in ws[harf][1:200] if x.value is not None])
            ws.column_dimensions[harf].width = min(max(10, uz + 2), 60)
            if any(isinstance(x.value, float) for x in ws[harf][1:50]):
                for x in ws[harf][1:]:
                    if isinstance(x.value, (int, float)):
                        x.number_format = "#,##0.00"
        ws.freeze_panes = "A2"
        if ws.max_row > 1 and cols:
            ws.auto_filter.ref = ws.dimensions
    if not wb.sheetnames:
        wb.create_sheet("Özet").append(["Veri bulunamadı"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_EXCEL_KAYNAKLARI = ["donem_ozeti", "kanal_tablosu", "zaman_serisi", "urun_raporu", "kategori_raporu",
                     "odeme_yontemleri", "saat_gun_analizi", "il_ilce", "iade_nedenleri",
                     "iptal_iade_urunleri", "stok_durumu", "hizli_satanlar", "kupon_performansi",
                     "donemin_urun_beden_dokumu", "siparis_listesi"]


async def _t_excel(a, u):
    """Seçilen rapor aracını TAM liste modunda çalıştırıp .xlsx üretir. Rakamlar dil
    modelinden DEĞİL, doğrudan rapor hesaplarından gelir."""
    kaynak = str(a.get("rapor") or "").strip()
    if kaynak not in _EXCEL_KAYNAKLARI:
        return {"hata": f"rapor şunlardan biri olmalı: {', '.join(_EXCEL_KAYNAKLARI)}"}
    args = {k: v for k, v in (a or {}).items() if k not in ("rapor", "dosya_adi")}
    args["_excel"] = True
    veri = _temizle(await ARACLAR[kaynak][0](args, u))
    if isinstance(veri, dict) and veri.get("hata"):
        return veri
    sayfalar = _sayfalara_bol(veri if isinstance(veri, dict) else {"veri": veri})
    ad = re.sub(r"[^\w\- ]+", "", str(a.get("dosya_adi") or kaynak), flags=re.U).strip()[:60] or kaynak
    s, e = (veri.get("aralik") or ["", ""]) if isinstance(veri, dict) else ("", "")
    dosya = f"{ad}{('_' + str(s) + '_' + str(e)) if s else ''}.xlsx".replace(" ", "_")
    icerik = _xlsx(ad, sayfalar)
    fid = uuid.uuid4().hex
    from bson import Binary
    await db.report_assistant_files.insert_one({
        "id": fid, "name": dosya, "data": Binary(icerik), "size": len(icerik),
        "user": u.get("email") or u.get("id"), "created_at": datetime.now(timezone.utc).isoformat()})
    liste = _DOSYALAR.get()
    if liste is not None:
        liste.append({"id": fid, "name": dosya, "size": len(icerik)})
    return {"dosya_hazir": True, "dosya_adi": dosya,
            "sayfalar": [{"ad": n, "satir": len(r)} for n, r in sayfalar],
            "not": "Dosya sohbette 'Excel'i indir' düğmesiyle kullanıcıya sunuldu."}


_T = lambda: {"type": "string", "description": "YYYY-MM-DD (Türkiye günü). Verilmezse ayın 1'i / bugün."}
_K = {"type": "string", "enum": _KANALLAR, "description": "Kanal; 'all' = tüm kanallar."}

ARACLAR: dict[str, tuple] = {
    "donem_ozeti": (_t_donem_ozeti,
        "Bir dönemin ciro özeti: brüt ciro, iptal, iade, DÖNEM (muhasebe) net'i, BU DÖNEMİN "
        "SATIŞININ net'i (kohort — pazaryeri panelleriyle kıyaslanan), önceki aylara ait iade/"
        "iptal, açık iade talepleri. Ciro/net/iade/iptal sorularının çoğu için ilk araç.",
        {"start_date": _T(), "end_date": _T(), "kanal": _K}),
    "kanal_tablosu": (_t_kanal_tablosu,
        "Kanal kanal (Site, Trendyol, Hepsiburada, Amazon, Temu) brüt/iptal/iade/net ve açık iade.",
        {"start_date": _T(), "end_date": _T()}),
    "zaman_serisi": (_t_zaman_serisi,
        "Gün gün, hafta hafta ya da ay ay net ciro ve sipariş serisi. Trend ve kıyas soruları için.",
        {"start_date": _T(), "end_date": _T(), "kanal": _K,
         "gruplama": {"type": "string", "enum": ["day", "week", "month"]}}),
    "urun_raporu": (_t_urun_raporu,
        "Ürün ürün satış: adet, ciro, iptal, iade, stok, beden ve kanal kırılımı. 'arama' ile ürün "
        "adı/stok kodu/barkod süzülür; 'siralama' ciro|adet|iade|iptal|stok.",
        {"start_date": _T(), "end_date": _T(), "kanal": _K,
         "arama": {"type": "string"},
         "siralama": {"type": "string", "enum": ["ciro", "adet", "iade", "iptal", "stok"]},
         "limit": {"type": "integer"}}),
    "kategori_raporu": (_t_kategori, "Kategori kategori satış adedi ve ciro.",
        {"start_date": _T(), "end_date": _T(), "kanal": _K}),
    "odeme_yontemleri": (_t_odeme, "Ödeme yöntemine göre (kart, havale, kapıda, pazaryeri) sipariş ve ciro.",
        {"start_date": _T(), "end_date": _T(), "kanal": _K}),
    "saat_gun_analizi": (_t_saat_gun,
        "Satışların saatlere (TR saati) ya da haftanın günlerine dağılımı ve zirve.",
        {"start_date": _T(), "end_date": _T(), "kanal": _K,
         "tur": {"type": "string", "enum": ["saat", "gun"]}}),
    "il_ilce": (_t_konum, "İl ya da ilçe bazında sipariş ve ciro (toplu).",
        {"start_date": _T(), "end_date": _T(), "kanal": _K,
         "seviye": {"type": "string", "enum": ["il", "ilce"]}, "limit": {"type": "integer"}}),
    "iade_nedenleri": (_t_iade_nedenleri, "İade nedenlerinin dağılımı.",
        {"start_date": _T(), "end_date": _T(), "kanal": _K}),
    "iptal_iade_urunleri": (_t_iptal_iade_urunleri,
        "En çok iade ya da iptal edilen ürünler.",
        {"start_date": _T(), "end_date": _T(), "kanal": _K,
         "tur": {"type": "string", "enum": ["iade", "iptal"]}, "limit": {"type": "integer"}}),
    "stok_durumu": (_t_stok, "Anlık stok: toplam adet/değer, azalan ve tükenen ürünler.", {}),
    "hizli_satanlar": (_t_hizli_satan, "Son N günde hızlı satan ürünler.",
        {"gun": {"type": "integer"}, "limit": {"type": "integer"}}),
    "kupon_performansi": (_t_kupon, "Kupon kupon kullanım, indirim ve ciro.",
        {"start_date": _T(), "end_date": _T()}),
    "donemin_urun_beden_dokumu": (_t_gunun_urunleri,
        "Bir gün ya da aralıkta sipariş edilen ürün+beden adetleri ve sipariş sayısı "
        "('dün ne sattık', 'bugün hangi bedenler gitti').",
        {"start_date": _T(), "end_date": _T(), "kanal": _K}),
    "siparis_listesi": (_t_siparis_listesi,
        "Dönemin sipariş listesi (sipariş no, kanal, TR tarih, statü, tutar, adet, ürünler). "
        "Sohbette en fazla 50 satır; tam liste için excel_olustur(rapor='siparis_listesi'). "
        "statu: 'iptal' yalnız iptaller, 'aktif' iptal hariç, 'hepsi' = tümü.",
        {"start_date": _T(), "end_date": _T(), "kanal": _K,
         "statu": {"type": "string", "enum": ["hepsi", "iptal", "aktif"]}, "limit": {"type": "integer"}}),
    "excel_olustur": (_t_excel,
        "İndirilebilir EXCEL (.xlsx) dosyası üretir. Kullanıcı Excel/liste/dosya/tablo indirmek "
        "istediğinde MUTLAKA bunu çağır: 'rapor' = hangi veri (ör. siparis_listesi, urun_raporu, "
        "kanal_tablosu, zaman_serisi...), diğer alanlar o raporun parametreleri. Dosya TAM liste "
        "içerir ve sohbette indirme düğmesi olarak görünür.",
        {"rapor": {"type": "string", "enum": _EXCEL_KAYNAKLARI},
         "dosya_adi": {"type": "string", "description": "Kısa dosya adı (Türkçe olabilir)"},
         "start_date": _T(), "end_date": _T(), "kanal": _K,
         "gruplama": {"type": "string", "enum": ["day", "week", "month"]},
         "siralama": {"type": "string", "enum": ["ciro", "adet", "iade", "iptal", "stok"]},
         "arama": {"type": "string"},
         "tur": {"type": "string", "description": "saat|gun (saat_gun_analizi) ya da iade|iptal"},
         "seviye": {"type": "string", "enum": ["il", "ilce"]},
         "statu": {"type": "string", "enum": ["hepsi", "iptal", "aktif"]},
         "gun": {"type": "integer"}}),
    "siparis_sorgula": (_t_siparis,
        "Tek siparişin durumu: statü, tarih, tutar, kalemler, kargo, gider pusulası, iade talebi. "
        "Müşteri kişisel bilgisi döndürmez.",
        {"siparis_no": {"type": "string"}}),
}


# ── Gizlilik süzgeci (son savunma hattı) ──────────────────────────────────────
_KISISEL = re.compile(r"(phone|telefon|email|e-posta|address|adres|first_name|last_name|"
                      r"full_name|customer_name|musteri|müşteri_adi|tc_no|tckn|iban|alici|alıcı)",
                      re.I)
_EPOSTA = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
# Sınırlarda rakam OLMAMALI: 11 haneli sipariş numaraları ve barkodlar telefon sanılmasın.
_TEL = re.compile(r"(?<!\d)(?:\+?90\s?|0)?5\d{2}\s?\d{3}\s?\d{2}\s?\d{2}(?!\d)")


def _temizle(o):
    if isinstance(o, dict):
        return {k: _temizle(v) for k, v in o.items() if not _KISISEL.search(str(k))}
    if isinstance(o, list):
        return [_temizle(v) for v in o]
    if isinstance(o, str):
        return _TEL.sub("[tel]", _EPOSTA.sub("[e-posta]", o))
    if isinstance(o, float) and (o != o or o in (float("inf"), float("-inf"))):
        return None
    return o


async def _arac_calistir(ad: str, args: dict, user: dict) -> str:
    t = ARACLAR.get(ad)
    if not t:
        return json.dumps({"hata": f"bilinmeyen araç: {ad}"}, ensure_ascii=False)
    try:
        sonuc = _temizle(await t[0](dict(args or {}), user))
    except HTTPException as he:
        sonuc = {"hata": str(he.detail)[:300]}
    except Exception as e:
        logger.warning(f"[rapor-asistan] {ad} hata: {e}")
        sonuc = {"hata": f"{type(e).__name__}: {str(e)[:200]}"}
    s = json.dumps(sonuc, ensure_ascii=False, default=str)
    if len(s) > _MAX_SONUC:
        s = s[:_MAX_SONUC] + '… [KISALTILDI — daha dar aralık ya da limit iste]'
    return s


def _sistem_mesaji() -> str:
    n = _tr_simdi()
    bu_ay = n.date().replace(day=1)
    gecen_ay_son = bu_ay - timedelta(days=1)
    return f"""Sen mağazanın rapor asistanısın. Yöneticinin satış/sipariş/iade/iptal/ürün/stok
sorularını, sistemin KENDİ rapor araçlarını çağırarak Türkçe cevaplarsın.

Şu an (Türkiye saati): {n.strftime('%d.%m.%Y %H:%M, %A')}.
Bu ay: {bu_ay.isoformat()} → {n.date().isoformat()}. Geçen ay: {gecen_ay_son.replace(day=1).isoformat()} → {gecen_ay_son.isoformat()}.

KESİN KURALLAR
- Hiçbir rakamı UYDURMA ve tahmin etme. Her sayı bir araç sonucundan gelmeli. Araç veri
  vermediyse bunu açıkça söyle.
- Soru dönem içeriyorsa tarihleri kendin hesapla (dün, geçen hafta, geçen ay, son 30 gün,
  "haziran" = bu yılın haziranı) ve cevapta hangi aralığı kullandığını yaz.
- Kıyas sorularında (bu ay vs geçen ay, kanal kanal) her dönem için aracı ayrı çağır.
- Müşterinin adı, telefonu, adresi gibi kişisel bilgileri isteseler bile verme.

SİSTEMİN SAYMA KURALLARI (cevaplarken bunlara göre açıkla)
- Satış, siparişin VERİLDİĞİ güne (TR saati) yazılır.
- İptal, iptalin kesinleştiği güne; İADE yalnız koçan numarası atanmış GİDER PUSULASI
  kesildiğinde ve pusulanın tarihine yazılır. Açık/onaylanmamış iade talebi satış sayılır.
- İki net vardır: "net_BU_DONEMIN_SATISININ" (kohort: bu dönemde satılanın net'i —
  Trendyol/HB panelleriyle kıyaslanacak olan, ticari performans) ve "net_DONEM_muhasebe"
  (bu dönemde kesinleşen TÜM iptal/iadeler düşülür, önceki ayların siparişlerinden gelenler
  dahil). "Net satış ne?" sorusuna varsayılan olarak kohort net'i ver, muhasebe net'ini
  ayrıca belirt.
- Trendyol panelindeki sipariş sayısı iptallerin bir kısmını saymaz; "Siparişleriniz"
  Excel'i iptalleri hiç içermez. Bu yüzden panel sayısı bizden biraz düşük olabilir.
- Kanal anahtarları: site, trendyol, hepsiburada, amazon, temu.

EXCEL / DOSYA
- Kullanıcı Excel, liste, dosya, indir, tablo olarak ver derse excel_olustur aracını çağır.
  Dosya sohbette "Excel'i indir" düğmesi olarak otomatik görünür. ASLA CSV/TSV metni ya da
  Python kodu önerme, "dosya gönderemem" deme. Cevapta kısaca dosyada ne olduğunu yaz.

BİÇİM
- Kısa ve net ol; önce cevabı ver, sonra gerekiyorsa açıkla.
- Tutarları Türk biçiminde yaz: 1.234.567,89 TL. Birden çok satır varsa küçük bir markdown
  tablo kullan.
"""


# ── Sağlayıcılar ──────────────────────────────────────────────────────────────
def _json_sema(ozellikler: dict) -> dict:
    return {"type": "object", "properties": ozellikler or {}}


async def _gemini(key: str, model: str, gecmis: list, user: dict, kullanilan: list) -> str:
    from google import genai
    from google.genai import types
    client = genai.Client(api_key=key)
    decls = [types.FunctionDeclaration(name=ad, description=t[1],
                                       parameters_json_schema=_json_sema(t[2]))
             for ad, t in ARACLAR.items()]
    config = types.GenerateContentConfig(
        system_instruction=_sistem_mesaji(),
        tools=[types.Tool(function_declarations=decls)],
        automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True),
        max_output_tokens=2048, temperature=0.2)
    contents = [types.Content(role="model" if m["role"] == "assistant" else "user",
                              parts=[types.Part(text=m["text"])]) for m in gecmis]
    for _ in range(_MAX_TUR):
        resp = await client.aio.models.generate_content(model=model, contents=contents, config=config)
        calls = resp.function_calls or []
        if not calls:
            return (resp.text or "").strip()
        contents.append(resp.candidates[0].content)
        parts = []
        for fc in calls:
            kullanilan.append(fc.name)
            sonuc = await _arac_calistir(fc.name, dict(fc.args or {}), user)
            parts.append(types.Part.from_function_response(name=fc.name, response={"sonuc": sonuc}))
        contents.append(types.Content(role="user", parts=parts))
    return "Soru çok fazla adım gerektirdi; lütfen daha dar bir soru sorun."


async def _openai(key: str, model: str, gecmis: list, user: dict, kullanilan: list) -> str:
    from openai import AsyncOpenAI
    client = AsyncOpenAI(api_key=key)
    tools = [{"type": "function", "function": {"name": ad, "description": t[1],
                                               "parameters": _json_sema(t[2])}}
             for ad, t in ARACLAR.items()]
    msgs: list = [{"role": "system", "content": _sistem_mesaji()}]
    msgs += [{"role": m["role"], "content": m["text"]} for m in gecmis]
    for _ in range(_MAX_TUR):
        resp = await client.chat.completions.create(model=model, messages=msgs, tools=tools,
                                                    max_completion_tokens=2048)
        msg = resp.choices[0].message
        if not msg.tool_calls:
            return (msg.content or "").strip()
        msgs.append({"role": "assistant", "content": msg.content or "",
                     "tool_calls": [tc.model_dump() for tc in msg.tool_calls]})
        for tc in msg.tool_calls:
            kullanilan.append(tc.function.name)
            try:
                args = json.loads(tc.function.arguments or "{}")
            except Exception:
                args = {}
            msgs.append({"role": "tool", "tool_call_id": tc.id,
                         "content": await _arac_calistir(tc.function.name, args, user)})
    return "Soru çok fazla adım gerektirdi; lütfen daha dar bir soru sorun."


def _gecmisi_temizle(raw: Any) -> list:
    out = []
    for m in (raw or [])[-12:]:
        if not isinstance(m, dict):
            continue
        rol = "assistant" if m.get("role") == "assistant" else "user"
        t = str(m.get("text") or m.get("content") or "").strip()[:3000]
        if t:
            out.append({"role": rol, "text": t})
    # Sohbet kullanıcı mesajıyla bitmeli
    while out and out[-1]["role"] != "user":
        out.pop()
    return out


# ── Uçlar ─────────────────────────────────────────────────────────────────────
@router.post("/chat")
@(limiter.limit("20/minute") if limiter else (lambda f: f))
async def assistant_chat(request: Request, payload: dict,
                         current_user: dict = Depends(require_admin)):
    ayar = await _ayarlar()
    if not ayar["enabled"]:
        raise HTTPException(status_code=403, detail="Rapor asistanı kapalı")
    gecmis = _gecmisi_temizle(payload.get("messages"))
    if not gecmis:
        raise HTTPException(status_code=400, detail="Soru boş")
    try:  # raporların ticimax-kopya elemesi için önbellek (rapor router'ı bunu kendisi yapar)
        from .report_dedup import load_dup_order_numbers
        await load_dup_order_numbers()
    except Exception:
        pass

    t0 = time.monotonic()
    dosyalar: list = []
    _DOSYALAR.set(dosyalar)
    kullanilan: list = []
    denemeler = []
    cevap, saglayici = "", ""
    gk = await _anahtar("GEMINI_API_KEY", "gemini|google")
    if gk:
        try:
            cevap = await _gemini(gk, ayar["gemini_model"], gecmis, current_user, kullanilan)
            saglayici = "Gemini"
        except Exception as e:
            denemeler.append(f"gemini: {type(e).__name__}: {str(e)[:160]}")
            logger.warning(f"[rapor-asistan] Gemini başarısız, OpenAI'ye geçiliyor: {e}")
    if not cevap:
        ok_ = await _anahtar("OPENAI_API_KEY", "openai")
        if ok_:
            try:
                kullanilan.clear()
                cevap = await _openai(ok_, ayar["openai_model"], gecmis, current_user, kullanilan)
                saglayici = "ChatGPT"
            except Exception as e:
                denemeler.append(f"openai: {type(e).__name__}: {str(e)[:160]}")
                logger.warning(f"[rapor-asistan] OpenAI başarısız: {e}")
    if not cevap:
        if not gk and not denemeler:
            raise HTTPException(status_code=503, detail="Gemini/OpenAI anahtarı tanımlı değil")
        raise HTTPException(status_code=502, detail="Asistan şu an cevap veremedi: " + " | ".join(denemeler)[:300])

    try:
        await db.report_assistant_logs.insert_one({
            "at": datetime.now(timezone.utc).isoformat(),
            "user": current_user.get("email") or current_user.get("id"),
            "soru": gecmis[-1]["text"][:2000], "cevap": cevap[:6000],
            "saglayici": saglayici, "araclar": kullanilan[:30],
            "ms": int((time.monotonic() - t0) * 1000), "denemeler": denemeler,
        })
    except Exception:
        pass
    return {"answer": cevap, "provider": saglayici, "tools": kullanilan,
            "files": [{"id": f["id"], "name": f["name"], "size": f["size"]} for f in dosyalar]}


@router.get("/file/{file_id}")
async def assistant_file(file_id: str, current_user: dict = Depends(require_admin)):
    """Asistanın ürettiği Excel'i indirir (yalnız admin; 7 gün saklanır)."""
    if not re.fullmatch(r"[0-9a-f]{32}", file_id or ""):
        raise HTTPException(status_code=404, detail="Dosya bulunamadı")
    d = await db.report_assistant_files.find_one({"id": file_id}, {"_id": 0})
    if not d:
        raise HTTPException(status_code=404, detail="Dosya bulunamadı")
    try:
        if datetime.fromisoformat(d["created_at"]) < datetime.now(timezone.utc) - timedelta(days=7):
            raise HTTPException(status_code=410, detail="Dosyanın süresi doldu, yeniden isteyin")
    except HTTPException:
        raise
    except Exception:
        pass
    from urllib.parse import quote
    ad = d.get("name") or "rapor.xlsx"
    return Response(content=bytes(d["data"]),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=\"rapor.xlsx\"; filename*=UTF-8''{quote(ad)}"})


@router.post("/transcribe")
@(limiter.limit("20/minute") if limiter else (lambda f: f))
async def assistant_transcribe(request: Request, audio: UploadFile = File(...),
                               current_user: dict = Depends(require_admin)):
    """Sesli soruyu yazıya çevirir (önce Gemini, olmazsa OpenAI Whisper)."""
    data = await audio.read()
    if not data:
        raise HTTPException(status_code=400, detail="Ses kaydı boş")
    if len(data) > _MAX_SES:
        raise HTTPException(status_code=413, detail="Ses kaydı çok uzun (en fazla ~10 MB)")
    mime = (audio.content_type or "audio/webm").split(";")[0].strip().lower()
    ayar = await _ayarlar()
    hatalar = []
    gk = await _anahtar("GEMINI_API_KEY", "gemini|google")
    if gk:
        try:
            import base64
            from google import genai
            client = genai.Client(api_key=gk)
            resp = await client.aio.models.generate_content(
                model=ayar["gemini_model"],
                contents=[{"role": "user", "parts": [
                    {"text": "Bu ses kaydını birebir Türkçe metne dök. SADECE konuşulan metni "
                             "yaz; yorum ya da açıklama ekleme."},
                    {"inline_data": {"mime_type": mime,
                                     "data": base64.standard_b64encode(data).decode()}}]}])
            txt = (getattr(resp, "text", None) or "").strip()
            if txt:
                return {"text": txt, "provider": "Gemini"}
        except Exception as e:
            hatalar.append(f"gemini: {type(e).__name__}")
            logger.warning(f"[rapor-asistan] Gemini ses çevirisi başarısız: {e}")
    ok_ = await _anahtar("OPENAI_API_KEY", "openai")
    if ok_:
        try:
            from openai import AsyncOpenAI
            ext = {"audio/mp4": "m4a", "audio/aac": "m4a", "audio/mpeg": "mp3",
                   "audio/ogg": "ogg", "audio/wav": "wav"}.get(mime, "webm")
            tr = await AsyncOpenAI(api_key=ok_).audio.transcriptions.create(
                model="whisper-1", file=(f"soru.{ext}", data, mime), language="tr")
            txt = (getattr(tr, "text", "") or "").strip()
            if txt:
                return {"text": txt, "provider": "Whisper"}
        except Exception as e:
            hatalar.append(f"openai: {type(e).__name__}")
            logger.warning(f"[rapor-asistan] Whisper başarısız: {e}")
    raise HTTPException(status_code=502, detail="Ses yazıya çevrilemedi" + (f" ({', '.join(hatalar)})" if hatalar else ""))
