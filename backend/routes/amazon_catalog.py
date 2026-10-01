"""
Amazon KATALOG EŞLEŞTİRME — Amazon'da (bu platform dışında) açılmış mevcut listelemeleri mağaza
ürün/varyantlarıyla eşleştirir, eşleşen SKU'ların stok/fiyatını senkrona dahil eder ve
"hangi ürünümüz Amazon'da var / yok" raporu üretir (JSON + Excel).

Akış:
  POST /amazon/spapi/catalog/sync      → arka planda: Reports API GET_MERCHANT_LISTINGS_ALL_DATA
                                          (tüm SellerSKU/ASIN/EAN/fiyat/stok) → db.amazon_catalog;
                                          eşleştirme → db.amazon_sku_map (matched / unmatched)
  GET  /amazon/spapi/catalog/status    → ilerleme + sayaçlar
  GET  /amazon/spapi/catalog/report    → özet + listeler (JSON)
  GET  /amazon/spapi/catalog/report.xlsx
  POST /amazon/spapi/catalog/push      → eşleşen (harici) SKU'lara stok/fiyat gönder (manuel, force)

Eşleştirme sırası (SKU başına ilk tutan): (1) Amazon product-id (EAN/barkod) = varyant barkodu,
(2) SellerSKU = bizim güncel SKU'muz (stokkodu-renk-beden), (3) SellerSKU = eski biçim
(stokkodu-beden), (4) SellerSKU = barkod, (5) SellerSKU = stok kodu (tek varyantlı ürün).
"""
import asyncio
import csv
import gzip
import io
import re
from datetime import datetime, timezone
from typing import Optional

import httpx
from fastapi import APIRouter, Depends, HTTPException, Query, Body
from fastapi.responses import StreamingResponse

from .deps import db, logger, require_admin

router = APIRouter(prefix="/amazon/spapi/catalog", tags=["Amazon Catalog"])

_STATUS_ID = "amazon_catalog_sync"
_TASKS = set()

_HEADER_ALIASES = {
    "sku": ("seller-sku", "seller_sku", "sku", "satici-sku", "satıcı-sku", "satıcı sku"),
    "asin": ("asin1", "asin", "asin-1"),
    "name": ("item-name", "item_name", "urun-adi", "ürün-adı", "ürün adı", "title"),
    "price": ("price", "fiyat"),
    "qty": ("quantity", "adet", "miktar"),
    "pid": ("product-id", "product_id", "urun-kimligi", "ürün-kimliği", "ürün kimliği"),
    "pid_type": ("product-id-type", "product_id_type", "urun-kimligi-turu", "ürün-kimliği-türü"),
    "status": ("status", "durum"),
    "fc": ("fulfillment-channel", "fulfilment-channel", "fulfillment_channel"),
    "image": ("image-url", "image_url"),
    "open_date": ("open-date", "open_date"),
}


def _norm_h(h: str) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower().replace("_", "-"))


def _col_index(headers: list) -> dict:
    hn = [_norm_h(h) for h in headers]
    idx = {}
    for key, aliases in _HEADER_ALIASES.items():
        for a in aliases:
            a2 = _norm_h(a)
            if a2 in hn:
                idx[key] = hn.index(a2)
                break
    # Standart (İngilizce) rapor düzeni için konumsal yedek
    _std = ["item-name", "item-description", "listing-id", "seller-sku", "price", "quantity", "open-date",
            "image-url", "item-is-marketplace", "product-id-type", "zshop-shipping-fee", "item-note",
            "item-condition", "zshop-category1", "zshop-browse-path", "zshop-storefront-feature", "asin1",
            "asin2", "asin3", "will-ship-internationally", "expedited-shipping", "zshop-boldface",
            "product-id", "bid-for-featured-placement", "add-delete", "pending-quantity",
            "fulfillment-channel", "merchant-shipping-group", "status"]
    if "sku" not in idx and len(headers) >= 4:
        for key, aliases in _HEADER_ALIASES.items():
            for a in aliases:
                if a in _std and _std.index(a) < len(headers):
                    idx.setdefault(key, _std.index(a))
                    break
    return idx


async def _write_status(**kw):
    doc = {"id": _STATUS_ID, "updated_at": datetime.now(timezone.utc).isoformat()}
    doc.update(kw)
    await db.settings.update_one({"id": _STATUS_ID}, {"$set": doc}, upsert=True)


async def _spapi_raw(method: str, path: str, body: dict = None, params: dict = None) -> dict:
    """Reports API çağrısı — SALT-OKUNUR iş (rapor isteme/okuma) olduğu için AMAZON_ALLOW_WRITE
    kapısından geçmez (o kapı listeleme/stok YAZMA çağrıları içindir)."""
    from .amazon_spapi import get_valid_access_token, _assert_restricted_allowed, _log_spapi_call, _scrub_pii
    _assert_restricted_allowed(path)
    token, endpoint, _ = await get_valid_access_token()
    async with httpx.AsyncClient(timeout=30.0) as c:
        r = await c.request(method, f"{endpoint}{path}", params=params or {}, json=body,
                            headers={"x-amz-access-token": token, "content-type": "application/json",
                                     "accept": "application/json"})
        try:
            data = r.json()
        except Exception:
            data = {"raw": r.text[:500]}
        data = _scrub_pii(data)
        ok = 200 <= r.status_code < 300
        await _log_spapi_call(f"{method} {path}", r.status_code, ok)
        return {"status": r.status_code, "ok": ok, "data": data}


async def fetch_amazon_report_rows(report_type: str, extra: dict = None, timeout_s: int = 300) -> list:
    """Genel Reports API akışı: rapor iste → hazırlanmasını bekle → indir → tab ayrılmış tabloyu
    {normalize-başlık: değer} sözlüklerine çevir. (İade raporu vb. için; listeleme raporu ayrı.)"""
    from .amazon_spapi import get_valid_access_token
    _, _, mp = await get_valid_access_token()
    body = {"reportType": report_type, "marketplaceIds": [mp]}
    body.update(extra or {})
    cr = await _spapi_raw("POST", "/reports/2021-06-30/reports", body)
    if not cr["ok"]:
        raise RuntimeError(f"Rapor isteği reddedildi (HTTP {cr['status']}): {str(cr['data'])[:300]}")
    rid = (cr["data"] or {}).get("reportId")
    if not rid:
        raise RuntimeError(f"reportId yok: {str(cr['data'])[:200]}")
    doc_id, waited = None, 0
    while waited < timeout_s:
        await asyncio.sleep(10)
        waited += 10
        st = await _spapi_raw("GET", f"/reports/2021-06-30/reports/{rid}")
        if not st["ok"]:
            continue
        ps = (st["data"] or {}).get("processingStatus")
        if ps == "DONE":
            doc_id = (st["data"] or {}).get("reportDocumentId")
            break
        if ps in ("CANCELLED", "FATAL"):
            # CANCELLED: seçilen aralıkta veri yok (Amazon böyle döner) → boş liste
            if ps == "CANCELLED":
                return []
            raise RuntimeError(f"Amazon raporu {ps} döndü")
    if not doc_id:
        raise RuntimeError("Rapor zamanında hazırlanmadı")
    dd = await _spapi_raw("GET", f"/reports/2021-06-30/documents/{doc_id}")
    if not dd["ok"]:
        raise RuntimeError(f"Rapor belgesi alınamadı (HTTP {dd['status']})")
    url = (dd["data"] or {}).get("url")
    comp = (dd["data"] or {}).get("compressionAlgorithm")
    async with httpx.AsyncClient(timeout=60.0, follow_redirects=True) as c:
        r = await c.get(url)
        r.raise_for_status()
        raw = r.content
    if comp == "GZIP":
        raw = gzip.decompress(raw)
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1254", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except Exception:
            continue
    text = text or raw.decode("latin-1", errors="ignore")
    rows = list(csv.reader(io.StringIO(text), delimiter="\t"))
    if not rows or len(rows) < 2:
        return []
    heads = [str(h or "").strip().lower().replace("_", "-").replace(" ", "-") for h in rows[0]]
    out = []
    for row in rows[1:]:
        if not any(str(x).strip() for x in row):
            continue
        out.append({heads[i]: (row[i] if i < len(row) else "") for i in range(len(heads))})
    return out


async def fetch_amazon_listings_report(timeout_s: int = 420) -> list:
    """GET_MERCHANT_LISTINGS_ALL_DATA raporunu ister, bekler, indirir, satırlara çevirir."""
    from .amazon_spapi import get_valid_access_token
    _, _, mp = await get_valid_access_token()
    cr = await _spapi_raw("POST", "/reports/2021-06-30/reports",
                          {"reportType": "GET_MERCHANT_LISTINGS_ALL_DATA", "marketplaceIds": [mp]})
    if not cr["ok"]:
        raise RuntimeError(f"Rapor isteği reddedildi (HTTP {cr['status']}): {str(cr['data'])[:300]}")
    rid = (cr["data"] or {}).get("reportId")
    if not rid:
        raise RuntimeError(f"reportId yok: {str(cr['data'])[:200]}")
    doc_id = None
    waited = 0
    while waited < timeout_s:
        await asyncio.sleep(12)
        waited += 12
        st = await _spapi_raw("GET", f"/reports/2021-06-30/reports/{rid}")
        if not st["ok"]:
            continue
        ps = (st["data"] or {}).get("processingStatus")
        if ps == "DONE":
            doc_id = (st["data"] or {}).get("reportDocumentId")
            break
        if ps in ("CANCELLED", "FATAL"):
            raise RuntimeError(f"Amazon raporu {ps} döndü")
        await _write_status(status="running", step=f"rapor hazırlanıyor ({ps}, {waited} sn)")
    if not doc_id:
        raise RuntimeError("Rapor zamanında hazırlanmadı (7 dk)")
    dd = await _spapi_raw("GET", f"/reports/2021-06-30/documents/{doc_id}")
    if not dd["ok"]:
        raise RuntimeError(f"Rapor belgesi alınamadı (HTTP {dd['status']})")
    url = (dd["data"] or {}).get("url")
    comp = (dd["data"] or {}).get("compressionAlgorithm")
    async with httpx.AsyncClient(timeout=60.0, follow_redirects=True) as c:
        r = await c.get(url)
        r.raise_for_status()
        raw = r.content
    if comp == "GZIP":
        raw = gzip.decompress(raw)
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1254", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except Exception:
            continue
    text = text or raw.decode("latin-1", errors="ignore")
    reader = csv.reader(io.StringIO(text), delimiter="\t")
    rows = list(reader)
    if not rows:
        return []
    idx = _col_index(rows[0])
    if "sku" not in idx:
        raise RuntimeError(f"Rapor başlıkları tanınamadı: {rows[0][:8]}")
    out = []

    def _g(row, key):
        i = idx.get(key)
        return (row[i].strip() if (i is not None and i < len(row)) else "")

    for row in rows[1:]:
        if not row or not _g(row, "sku"):
            continue
        try:
            qty = int(float(_g(row, "qty") or 0))
        except Exception:
            qty = 0
        try:
            price = float(str(_g(row, "price") or "0").replace(",", "."))
        except Exception:
            price = 0.0
        out.append({
            "sku": _g(row, "sku"), "asin": _g(row, "asin"), "item_name": _g(row, "name"),
            "price": price, "qty": qty, "product_id": _g(row, "pid"), "product_id_type": _g(row, "pid_type"),
            "status": _g(row, "status"), "fulfillment_channel": _g(row, "fc"),
            "image_url": _g(row, "image"), "open_date": _g(row, "open_date"),
        })
    return out


def _digits(s: str) -> str:
    return re.sub(r"\D", "", str(s or ""))


async def match_catalog(rows: list) -> dict:
    """Amazon satırlarını mağaza varyantlarıyla eşleştirir → db.amazon_sku_map."""
    from .amazon_spapi import _amazon_seller_sku
    from .integrations_common import _resolve_stock_code
    products = await db.products.find({"is_deleted": {"$ne": True}}, {"_id": 0, "id": 1, "name": 1, "stock_code": 1,
                                                                       "barcode": 1, "variants": 1, "is_active": 1,
                                                                       "price": 1, "sale_price": 1, "color": 1}).to_list(None)
    by_barcode, by_sku, by_legacy, by_code = {}, {}, {}, {}
    for p in products:
        code = _resolve_stock_code(p) or str(p.get("stock_code") or "").strip()
        vs = p.get("variants") or []
        if code and len(vs) == 1:
            by_code.setdefault(code, (p, vs[0]))
        for v in vs:
            bc = _digits(v.get("barcode"))
            if bc:
                by_barcode.setdefault(bc, (p, v))
            sk = _amazon_seller_sku(v, p)
            if sk:
                by_sku.setdefault(sk, (p, v))
            size = str(v.get("size") or "").strip()
            for c in {code, str(v.get("stock_code") or "").strip()} - {""}:
                if size:
                    by_legacy.setdefault(f"{c}-{size}", (p, v))
                    by_legacy.setdefault(f"{c}-{size}".upper(), (p, v))
    now = datetime.now(timezone.utc).isoformat()
    matched = unmatched = 0
    ops = []
    for r in rows:
        sku = r["sku"]
        hit, how = None, ""
        pid = _digits(r.get("product_id"))
        if pid and pid in by_barcode:
            hit, how = by_barcode[pid], "ean"
        elif sku in by_sku:
            hit, how = by_sku[sku], "sku"
        elif sku in by_legacy or sku.upper() in by_legacy:
            hit, how = by_legacy.get(sku) or by_legacy.get(sku.upper()), "legacy_sku"
        elif _digits(sku) and _digits(sku) in by_barcode and len(_digits(sku)) >= 8:
            hit, how = by_barcode[_digits(sku)], "sku_barcode"
        elif sku in by_code:
            hit, how = by_code[sku], "stock_code"
        doc = {"sku": sku, "asin": r.get("asin"), "item_name": r.get("item_name"),
               "amazon_price": r.get("price"), "amazon_qty": r.get("qty"), "amazon_status": r.get("status"),
               "amazon_product_id": r.get("product_id"), "fulfillment_channel": r.get("fulfillment_channel"),
               "matched": bool(hit), "match_by": how, "updated_at": now}
        if hit:
            p, v = hit
            doc.update({"product_id": p.get("id"), "product_name": p.get("name"),
                        "variant_barcode": str(v.get("barcode") or ""), "variant_size": str(v.get("size") or ""),
                        "variant_color": str(v.get("color") or p.get("color") or ""),
                        "our_sku": _amazon_seller_sku(v, p), "product_active": bool(p.get("is_active"))})
            matched += 1
        else:
            unmatched += 1
        ops.append(doc)
    await db.amazon_sku_map.delete_many({})
    if ops:
        await db.amazon_sku_map.insert_many(ops)
    return {"rows": len(rows), "matched": matched, "unmatched": unmatched}


async def _enrich_product_types(limit: int = 400):
    """Eşleşen HARİCİ SKU'lar için Amazon productType (PATCH stok/fiyat için gerekli) — Listings GET."""
    from .amazon_spapi import _spapi_get, _require_seller_id, get_valid_access_token
    seller = await _require_seller_id()
    _, _, mp = await get_valid_access_token()
    n = 0
    cur = db.amazon_sku_map.find({"matched": True, "$or": [{"product_type": {"$exists": False}}, {"product_type": ""}]},
                                 {"_id": 0, "sku": 1}).limit(limit)
    async for m in cur:
        try:
            r = await _spapi_get(f"/listings/2021-08-01/items/{seller}/{m['sku']}", {"marketplaceIds": mp, "includedData": "summaries"})
            if r["ok"]:
                s0 = ((r["data"] or {}).get("summaries") or [{}])[0]
                await db.amazon_sku_map.update_one({"sku": m["sku"]}, {"$set": {
                    "product_type": s0.get("productType") or "", "listing_status": s0.get("status"),
                    "asin": s0.get("asin") or None}})
                n += 1
        except Exception as e:
            logger.warning(f"[amazon-catalog] PT enrich {m['sku']}: {e}")
        await asyncio.sleep(0.25)
    return n


async def run_catalog_sync():
    started = datetime.now(timezone.utc)
    await _write_status(status="running", started_at=started.isoformat(), step="rapor isteniyor", error="")
    try:
        rows = await fetch_amazon_listings_report()
        await _write_status(status="running", step=f"{len(rows)} satır alındı, eşleştiriliyor")
        await db.amazon_catalog.delete_many({})
        if rows:
            await db.amazon_catalog.insert_many([dict(r, fetched_at=started.isoformat()) for r in rows])
        st = await match_catalog(rows)
        await _write_status(status="running", step="Amazon ürün tipleri çekiliyor", **st)
        n_pt = await _enrich_product_types()
        await _write_status(status="ok", finished_at=datetime.now(timezone.utc).isoformat(),
                            step="bitti", product_types_fetched=n_pt, **st)
        logger.info(f"[amazon-catalog] senkron bitti: {st} pt={n_pt}")
    except Exception as e:
        logger.exception(f"[amazon-catalog] senkron hatası: {e}")
        await _write_status(status="error", finished_at=datetime.now(timezone.utc).isoformat(),
                            error=str(e)[:400])


@router.post("/sync")
async def catalog_sync_start(current_user: dict = Depends(require_admin)):
    st = await db.settings.find_one({"id": _STATUS_ID}, {"_id": 0}) or {}
    if st.get("status") == "running":
        try:
            _dt = datetime.fromisoformat(st.get("started_at"))
            if (datetime.now(timezone.utc) - _dt).total_seconds() < 900:
                return {"success": True, "already_running": True, **st}
        except Exception:
            pass
    t = asyncio.create_task(run_catalog_sync())
    _TASKS.add(t)
    t.add_done_callback(_TASKS.discard)
    return {"success": True, "started": True,
            "message": "Amazon kataloğu arka planda çekiliyor (rapor 1-5 dk sürebilir). Durum: /catalog/status"}


@router.get("/status")
async def catalog_sync_status(current_user: dict = Depends(require_admin)):
    st = await db.settings.find_one({"id": _STATUS_ID}, {"_id": 0}) or {"status": "unknown"}
    # Sunucu yeniden başlarken kesilen çekim "running" olarak takılı kalıp butonu kilitliyordu.
    if st.get("status") == "running" and not any(not t.done() for t in _TASKS):
        try:
            if (datetime.now(timezone.utc) - datetime.fromisoformat(st.get("started_at"))).total_seconds() > 900:
                st["status"], st["step"] = "error", "yarıda kesildi — tekrar çekin"
        except Exception:
            pass
    st["map_matched"] = await db.amazon_sku_map.count_documents({"matched": True})
    st["map_unmatched"] = await db.amazon_sku_map.count_documents({"matched": False})
    return st


async def _build_report(limit: int = 2000) -> dict:
    from .amazon_spapi import _amazon_seller_sku, _amazon_markup, _amazon_price_of
    markup = await _amazon_markup()
    maps = await db.amazon_sku_map.find({"matched": True}, {"_id": 0}).to_list(None)
    by_variant = {}
    for m in maps:
        by_variant.setdefault((m.get("product_id"), m.get("variant_barcode") or m.get("variant_size")), []).append(m)
    products = await db.products.find({"is_active": True, "is_deleted": {"$ne": True}},
                                      {"_id": 0, "id": 1, "name": 1, "stock_code": 1, "variants": 1, "price": 1,
                                       "sale_price": 1, "member_price_1": 1, "color": 1}).sort("name", 1).to_list(None)
    on, off, partial = [], [], []
    variant_rows = []
    for p in products:
        vs = p.get("variants") or []
        our_price = _amazon_price_of(p, markup)
        n_on = 0
        for v in vs:
            key = (p.get("id"), str(v.get("barcode") or "") or str(v.get("size") or ""))
            ms = by_variant.get(key) or []
            if ms:
                n_on += 1
            for m in (ms or [None]):
                variant_rows.append({
                    "product": p.get("name"), "stock_code": p.get("stock_code") or "", "size": v.get("size") or "",
                    "color": v.get("color") or p.get("color") or "", "barcode": v.get("barcode") or "",
                    "our_stock": int(v.get("stock") or 0), "our_price": our_price,
                    "on_amazon": bool(m), "amazon_sku": (m or {}).get("sku") or "", "asin": (m or {}).get("asin") or "",
                    "match_by": (m or {}).get("match_by") or "", "amazon_qty": (m or {}).get("amazon_qty"),
                    "amazon_price": (m or {}).get("amazon_price"), "amazon_status": (m or {}).get("amazon_status") or "",
                    "product_type": (m or {}).get("product_type") or "",
                })
        row = {"product": p.get("name"), "stock_code": p.get("stock_code") or "", "product_id": p.get("id"),
               "variants": len(vs), "variants_on_amazon": n_on,
               "missing_sizes": [str(v.get("size") or "") for v in vs
                                 if not by_variant.get((p.get("id"), str(v.get("barcode") or "") or str(v.get("size") or "")))]}
        if n_on == 0:
            off.append(row)
        elif n_on < len(vs):
            partial.append(row)
        else:
            on.append(row)
    unmatched = await db.amazon_sku_map.find({"matched": False}, {"_id": 0}).sort("item_name", 1).to_list(limit)
    st = await db.settings.find_one({"id": _STATUS_ID}, {"_id": 0}) or {}
    return {
        "catalog_status": st,
        "summary": {"our_active_products": len(products), "on_amazon_full": len(on), "on_amazon_partial": len(partial),
                    "not_on_amazon": len(off), "amazon_rows": await db.amazon_catalog.count_documents({}),
                    "amazon_matched_skus": len(maps), "amazon_unmatched_skus": len(unmatched)},
        "on_amazon": on, "partial": partial, "not_on_amazon": off, "amazon_unmatched": unmatched,
        "variants": variant_rows,
    }


@router.get("/report")
async def catalog_report(current_user: dict = Depends(require_admin)):
    rep = await _build_report()
    rep["variants"] = rep["variants"][:3000]
    return rep


@router.get("/report.xlsx")
async def catalog_report_xlsx(current_user: dict = Depends(require_admin)):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    rep = await _build_report(limit=20000)
    wb = openpyxl.Workbook()

    def _sheet(ws, headers, rows, widths=None):
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="232F3E")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [18] * len(headers), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    s = rep["summary"]
    ws = wb.active
    ws.title = "Özet"
    _sheet(ws, ["Ölçüt", "Değer"], [
        ["Aktif ürünümüz", s["our_active_products"]],
        ["Amazon'da TÜM bedenleriyle olan", s["on_amazon_full"]],
        ["Amazon'da KISMEN olan (bazı bedenler yok)", s["on_amazon_partial"]],
        ["Amazon'da OLMAYAN", s["not_on_amazon"]],
        ["Amazon'daki toplam listeleme (SKU)", s["amazon_rows"]],
        ["Eşleşen Amazon SKU", s["amazon_matched_skus"]],
        ["Eşleşmeyen Amazon SKU (bizde karşılığı yok)", s["amazon_unmatched_skus"]],
        ["Katalog çekim zamanı", (rep["catalog_status"] or {}).get("finished_at") or ""],
    ], [44, 24])
    _sheet(wb.create_sheet("Amazon'da OLMAYAN"), ["Ürün", "Stok Kodu", "Beden sayısı"],
           [[r["product"], r["stock_code"], r["variants"]] for r in rep["not_on_amazon"]], [44, 18, 12])
    _sheet(wb.create_sheet("Kısmen olan"), ["Ürün", "Stok Kodu", "Beden", "Amazon'da", "Eksik bedenler"],
           [[r["product"], r["stock_code"], r["variants"], r["variants_on_amazon"], ", ".join(r["missing_sizes"])] for r in rep["partial"]],
           [44, 18, 8, 10, 30])
    _sheet(wb.create_sheet("Amazon'da olan"), ["Ürün", "Stok Kodu", "Beden sayısı"],
           [[r["product"], r["stock_code"], r["variants"]] for r in rep["on_amazon"]], [44, 18, 12])
    _sheet(wb.create_sheet("Varyant detayı"),
           ["Ürün", "Stok Kodu", "Beden", "Renk", "Barkod", "Bizim stok", "Bizim fiyat (marjlı)", "Amazon'da?",
            "Amazon SKU", "ASIN", "Eşleşme", "Amazon stok", "Amazon fiyat", "Amazon durum", "productType"],
           [[v["product"], v["stock_code"], v["size"], v["color"], v["barcode"], v["our_stock"], v["our_price"],
             "Evet" if v["on_amazon"] else "Hayır", v["amazon_sku"], v["asin"], v["match_by"], v["amazon_qty"],
             v["amazon_price"], v["amazon_status"], v["product_type"]] for v in rep["variants"]],
           [40, 16, 8, 14, 16, 10, 14, 10, 28, 14, 12, 11, 12, 12, 16])
    _sheet(wb.create_sheet("Amazon'da eşleşmeyen"),
           ["Amazon SKU", "ASIN", "Ürün adı (Amazon)", "EAN/ID", "Amazon stok", "Amazon fiyat", "Durum"],
           [[u.get("sku"), u.get("asin"), u.get("item_name"), u.get("amazon_product_id"), u.get("amazon_qty"),
             u.get("amazon_price"), u.get("amazon_status")] for u in rep["amazon_unmatched"]],
           [28, 14, 50, 16, 11, 12, 12])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"amazon-katalog-{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={fname}", "Cache-Control": "no-store"})


@router.post("/push")
async def catalog_push(payload: dict = Body(default={}), current_user: dict = Depends(require_admin)):
    """Eşleşen (bu platform dışında açılmış) Amazon SKU'larına stok + marjlı fiyat gönder (manuel, force)."""
    from scheduler import _run_amazon_auto_stock_sync
    res = await _run_amazon_auto_stock_sync(force=True, manual=True, mapped_only=not payload.get("include_ours"))
    return res
